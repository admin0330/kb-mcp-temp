"""Text extraction for the knowledge-base indexer.

Each extractor is intentionally defensive: a corrupt or password-protected file
should never crash indexing, it should raise ``ExtractionError`` so the caller
can log it and move on. Heavy / optional third-party libraries are imported
lazily inside each function so a missing dependency produces an actionable
message instead of an import-time failure for the whole server.
"""
from __future__ import annotations

from pathlib import Path
from typing import Callable, Dict

# Extensions we can read as plain UTF-8 (with latin-1 fallback).
TEXT_EXT = {
    ".md", ".markdown", ".txt", ".rst", ".org", ".tex", ".adoc",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".mjs", ".cjs",
    ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf",
    ".c", ".h", ".cc", ".cpp", ".hpp", ".cs", ".java", ".kt", ".swift",
    ".go", ".rs", ".rb", ".php", ".pl", ".lua", ".r", ".scala", ".sql",
    ".sh", ".bash", ".zsh", ".fish", ".ps1", ".bat",
    ".html", ".htm", ".xml", ".css", ".scss", ".vue", ".svelte",
    ".csv", ".tsv", ".log", ".env", ".gitignore", ".dockerignore",
}

# Binary document formats with dedicated extractors.
DOC_EXT = {".pdf", ".docx", ".pptx", ".xlsx", ".xlsm", ".epub"}

#: All extensions this module knows how to read.
SUPPORTED_EXT = TEXT_EXT | DOC_EXT


class ExtractionError(RuntimeError):
    """Raised when a supported file cannot be read."""


def is_supported(path: Path) -> bool:
    """Return ``True`` when :func:`extract_text` can handle ``path``."""
    return path.suffix.lower() in SUPPORTED_EXT


def extract_text(path: Path, max_chars: int = 2_000_000) -> str:
    """Return the best-effort plain-text content of ``path``.

    :param path: File to read.
    :param max_chars: Hard cap on returned characters (protects memory on huge
        PDFs / spreadsheets).
    :raises ExtractionError: When the file type is unsupported or unreadable.
    """
    ext = path.suffix.lower()
    handler = _DISPATCH.get(ext)
    if handler is None:
        if ext in TEXT_EXT:  # defensive; should be in _DISPATCH already
            handler = _read_text
        else:
            raise ExtractionError(f"unsupported file type: {ext or '<none>'}")
    try:
        text = handler(path, max_chars)
    except ExtractionError:
        raise
    except ModuleNotFoundError as exc:  # missing optional dependency
        raise ExtractionError(
            f"missing dependency for {ext}: {exc.name} (pip install it)"
        ) from exc
    except Exception as exc:  # noqa: BLE001 - we want a uniform failure type
        raise ExtractionError(f"failed to read {path.name}: {exc}") from exc
    return text[:max_chars].strip()


# --------------------------------------------------------------------------- #
# Individual format handlers. Signature: (path, max_chars) -> str
# --------------------------------------------------------------------------- #
def _read_text(path: Path, max_chars: int) -> str:
    data = path.read_bytes()[: max_chars * 4]  # bytes >= chars
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(path: Path, max_chars: int) -> str:
    from pypdf import PdfReader  # lazy import

    reader = PdfReader(str(path))
    parts: list[str] = []
    total = 0
    for page in reader.pages:
        chunk = page.extract_text() or ""
        parts.append(chunk)
        total += len(chunk)
        if total >= max_chars:
            break
    return "\n".join(parts)


def _docx(path: Path, max_chars: int) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pptx(path: Path, max_chars: int) -> str:
    from pptx import Presentation  # python-pptx

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"# Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        parts.append(line)
    return "\n".join(parts)


def _xlsx(path: Path, max_chars: int) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    total = 0
    for ws in wb.worksheets:
        parts.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                line = " | ".join(cells)
                parts.append(line)
                total += len(line)
                if total >= max_chars:
                    wb.close()
                    return "\n".join(parts)
    wb.close()
    return "\n".join(parts)


def _epub(path: Path, max_chars: int) -> str:
    import ebooklib
    from ebooklib import epub
    from bs4 import BeautifulSoup

    book = epub.read_epub(str(path))
    parts: list[str] = []
    total = 0
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        soup = BeautifulSoup(item.get_content(), "lxml")
        text = soup.get_text(separator="\n").strip()
        if text:
            parts.append(text)
            total += len(text)
            if total >= max_chars:
                break
    return "\n".join(parts)


_DISPATCH: Dict[str, Callable[[Path, int], str]] = {
    **{ext: _read_text for ext in TEXT_EXT},
    ".pdf": _pdf,
    ".docx": _docx,
    ".pptx": _pptx,
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".epub": _epub,
}
